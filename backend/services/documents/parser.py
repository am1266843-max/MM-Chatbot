import os
import logging
from typing import List, Dict, Any

logger = logging.getLogger(__name__)

class DocumentParser:
    """Extracts text and metadata from various file formats."""

    @staticmethod
    def parse_file(file_path: str, filename: str) -> List[Dict[str, Any]]:
        """
        Parses document and returns a list of pages/sections:
        [{"page_number": int, "section_title": str, "text": str}]
        """
        ext = os.path.splitext(filename)[1].lower()

        if ext == ".pdf":
            return DocumentParser._parse_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            return DocumentParser._parse_docx(file_path)
        elif ext in [".pptx", ".ppt"]:
            return DocumentParser._parse_pptx(file_path)
        elif ext in [".csv"]:
            return DocumentParser._parse_csv(file_path)
        elif ext in [".txt", ".md"]:
            return DocumentParser._parse_plain_text(file_path)
        else:
            raise ValueError(f"Unsupported file extension '{ext}'. Supported formats: PDF, DOCX, PPTX, CSV, TXT, MD.")

    @staticmethod
    def _parse_pdf(file_path: str) -> List[Dict[str, Any]]:
        import fitz  # PyMuPDF
        results = []
        try:
            doc = fitz.open(file_path)
            for page_idx in range(len(doc)):
                page = doc[page_idx]
                text = page.get_text()
                if text and text.strip():
                    results.append({
                        "page_number": page_idx + 1,
                        "section_title": f"Page {page_idx + 1}",
                        "text": text.strip()
                    })
            doc.close()
        except Exception as e:
            logger.error(f"Error parsing PDF '{file_path}': {e}")
            raise RuntimeError(f"Failed to extract text from PDF: {e}")

        if not results:
            results.append({
                "page_number": 1,
                "section_title": "Empty/Scanned PDF",
                "text": "Note: No extractable text found in this PDF (it may contain scanned images)."
            })
        return results

    @staticmethod
    def _parse_docx(file_path: str) -> List[Dict[str, Any]]:
        import docx
        results = []
        try:
            doc = docx.Document(file_path)
            paragraphs = []
            for p in doc.paragraphs:
                if p.text.strip():
                    paragraphs.append(p.text.strip())

            # Also extract text from tables
            for table in doc.tables:
                table_rows = []
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        table_rows.append(" | ".join(row_cells))
                if table_rows:
                    paragraphs.append("\n".join(table_rows))

            full_text = "\n\n".join(paragraphs)
            results.append({
                "page_number": 1,
                "section_title": "Document Body",
                "text": full_text
            })
        except Exception as e:
            logger.error(f"Error parsing DOCX '{file_path}': {e}")
            raise RuntimeError(f"Failed to extract text from DOCX: {e}")

        return results

    @staticmethod
    def _parse_pptx(file_path: str) -> List[Dict[str, Any]]:
        from pptx import Presentation
        results = []
        try:
            prs = Presentation(file_path)
            for slide_idx, slide in enumerate(prs.slides):
                slide_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_texts.append(shape.text.strip())
                if slide_texts:
                    results.append({
                        "page_number": slide_idx + 1,
                        "section_title": f"Slide {slide_idx + 1}",
                        "text": "\n".join(slide_texts)
                    })
        except Exception as e:
            logger.error(f"Error parsing PPTX '{file_path}': {e}")
            raise RuntimeError(f"Failed to extract text from PPTX: {e}")

        return results

    @staticmethod
    def _parse_csv(file_path: str) -> List[Dict[str, Any]]:
        import pandas as pd
        results = []
        try:
            df = pd.read_csv(file_path)
            summary = f"Columns: {', '.join(df.columns)}\nTotal Rows: {len(df)}\n\nPreview (First 100 rows):\n"
            summary += df.head(100).to_string(index=False)
            results.append({
                "page_number": 1,
                "section_title": "Tabular Data",
                "text": summary
            })
        except Exception as e:
            logger.error(f"Error parsing CSV '{file_path}': {e}")
            raise RuntimeError(f"Failed to extract text from CSV: {e}")

        return results

    @staticmethod
    def _parse_plain_text(file_path: str) -> List[Dict[str, Any]]:
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
            return [{
                "page_number": 1,
                "section_title": "Text File",
                "text": content.strip()
            }]
        except Exception as e:
            logger.error(f"Error reading text file '{file_path}': {e}")
            raise RuntimeError(f"Failed to read text file: {e}")

